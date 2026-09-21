"""Excel export for auditable labeled-chat analysis."""

from __future__ import annotations

from pathlib import Path
from tempfile import NamedTemporaryFile
from typing import Any, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from .analysis import AnalysisTables


_SHEET_NAMES = ["标签明细", "情绪趋势", "主题", "原话"]
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
_DATA_ALIGNMENT = Alignment(vertical="top")
_WRAPPED_ALIGNMENT = Alignment(vertical="top", wrap_text=True)
_PERCENT_FORMAT = "0.0%"


def write_analysis_workbook(tables: AnalysisTables, destination: Path) -> Path:
    """Write analysis tables to a verified temporary workbook, then replace ``destination``."""
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary_path: Path | None = None
    try:
        with NamedTemporaryFile(
            mode="wb", suffix=".tmp.xlsx", prefix=f".{destination.stem}.", dir=destination.parent, delete=False
        ) as temporary:
            temporary_path = Path(temporary.name)

        workbook = Workbook()
        detail = workbook.active
        detail.title = "标签明细"
        trends = workbook.create_sheet("情绪趋势")
        topics = workbook.create_sheet("主题")
        quotes = workbook.create_sheet("原话")

        _write_label_detail(detail, tables.label_rows)
        _write_sentiment_trends(trends, tables.sentiment_rows)
        _write_topics(topics, tables.topic_summary_rows, tables.topic_trend_rows)
        _write_quotes(quotes, tables.quote_rows)

        workbook.save(temporary_path)
        _verify_workbook(temporary_path)
        temporary_path.replace(destination)
        return destination
    finally:
        if temporary_path is not None and temporary_path.exists():
            temporary_path.unlink()


def _write_label_detail(sheet: Worksheet, rows: Iterable[dict[str, Any]]) -> None:
    headers = [
        "时间（秒）", "时间（毫秒）", "ISO 时间", "用户", "清洗文本", "原始弹幕", "情绪", "主题",
        "兴趣信号", "编码警告", "标注服务", "标注模型", "批次号", "消息 ID",
    ]
    _write_header(sheet, 1, headers)
    for row in rows:
        sheet.append([
            row.get("timestamp_seconds"), row.get("timestamp_ms"), row.get("timestamp_iso") or None,
            row.get("author", ""), row.get("text", ""), row.get("original_text", ""), row.get("sentiment", ""),
            row.get("topic", ""), _yes_no(row.get("interest_signal")), _yes_no(row.get("encoding_warning")),
            row.get("label_provider", ""), row.get("label_model") or None, row.get("batch_number"), row.get("message_id", ""),
        ])
    _finish_detail_sheet(sheet, headers, wrap_columns={5, 6}, widths=[12, 14, 25, 18, 35, 48, 12, 18, 12, 12, 18, 18, 12, 22])


def _write_sentiment_trends(sheet: Worksheet, rows: Iterable[dict[str, Any]]) -> None:
    rows = list(rows)
    headers = [
        "开始秒", "结束秒", "弹幕数", "独立用户数", "正面数量", "中性数量", "负面数量",
        "正面占比", "中性占比", "负面占比", "兴趣信号数",
    ]
    header_row = 16 if rows else 1
    _write_header(sheet, header_row, headers)
    for row in rows:
        sheet.append([
            row.get("start_seconds"), row.get("end_seconds"), row.get("message_count"), row.get("unique_authors"),
            row.get("positive"), row.get("neutral"), row.get("negative"), row.get("positive_share"),
            row.get("neutral_share"), row.get("negative_share"), row.get("interest_signals"),
        ])
    if rows:
        _add_sentiment_chart(sheet, header_row, len(rows))
    else:
        sheet["A2"] = "没有可分析的弹幕"
    _finish_detail_sheet(sheet, headers, header_row=header_row, percent_columns={8, 9, 10}, widths=[12] * len(headers))


def _write_topics(
    sheet: Worksheet, summary_rows: Iterable[dict[str, Any]], trend_rows: Iterable[dict[str, Any]]
) -> None:
    summary_rows = list(summary_rows)
    trend_rows = list(trend_rows)
    summary_headers = [
        "主题", "弹幕数量", "占比", "首次出现秒", "末次出现秒", "正面数量", "中性数量", "负面数量",
        "代表性原话 1", "代表性原话 2", "代表性原话 3",
    ]
    _write_header(sheet, 1, summary_headers)
    for row in summary_rows:
        sheet.append([
            row.get("topic", ""), row.get("message_count"), row.get("share"), row.get("first_seconds"),
            row.get("last_seconds"), row.get("positive"), row.get("neutral"), row.get("negative"),
            row.get("representative_quote_1", ""), row.get("representative_quote_2", ""),
            row.get("representative_quote_3", ""),
        ])
    if summary_rows:
        _add_topic_chart(sheet, len(summary_rows))
    else:
        sheet["A2"] = "没有可分析的弹幕"

    _finish_detail_sheet(
        sheet, summary_headers, percent_columns={3}, wrap_columns={9, 10, 11},
        widths=[20, 12, 12, 14, 14, 12, 12, 12, 42, 42, 42],
    )

    trend_header_row = max(5, len(summary_rows) + 4)
    sheet.cell(trend_header_row - 1, 1, "时间窗口 × 主题")
    trend_headers = ["开始秒", "结束秒", "主题", "弹幕数量", "主题在该窗口内的占比", "正面数量", "中性数量", "负面数量"]
    _write_header(sheet, trend_header_row, trend_headers)
    for row in trend_rows:
        sheet.append([
            row.get("start_seconds"), row.get("end_seconds"), row.get("topic", ""), row.get("message_count"),
            row.get("topic_share"), row.get("positive"), row.get("neutral"), row.get("negative"),
        ])
    for row in sheet.iter_rows(min_row=trend_header_row + 1, max_row=sheet.max_row, max_col=len(trend_headers)):
        for cell in row:
            cell.alignment = _DATA_ALIGNMENT
        row[4].number_format = _PERCENT_FORMAT


def _write_quotes(sheet: Worksheet, rows: Iterable[dict[str, Any]]) -> None:
    headers = ["时间（秒）", "用户", "原话", "主题", "情绪", "消息 ID"]
    _write_header(sheet, 1, headers)
    for row in rows:
        sheet.append([
            row.get("timestamp_seconds"), row.get("author", ""), row.get("original_text", ""), row.get("topic", ""),
            row.get("sentiment", ""), row.get("message_id", ""),
        ])
    _finish_detail_sheet(sheet, headers, wrap_columns={3}, widths=[12, 18, 55, 20, 12, 22])


def _write_header(sheet: Worksheet, row: int, headers: list[str]) -> None:
    for column, header in enumerate(headers, 1):
        cell = sheet.cell(row, column, header)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = _HEADER_ALIGNMENT


def _finish_detail_sheet(
    sheet: Worksheet,
    headers: list[str],
    *,
    header_row: int = 1,
    percent_columns: set[int] | None = None,
    wrap_columns: set[int] | None = None,
    widths: list[int] | None = None,
) -> None:
    percent_columns = percent_columns or set()
    wrap_columns = wrap_columns or set()
    final_row = max(header_row, sheet.max_row)
    sheet.auto_filter.ref = f"A{header_row}:{_column_letter(len(headers))}{final_row}"
    sheet.freeze_panes = f"A{header_row + 1}"
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=final_row, max_col=len(headers)):
        for index, cell in enumerate(row, 1):
            cell.alignment = _WRAPPED_ALIGNMENT if index in wrap_columns else _DATA_ALIGNMENT
            if index in percent_columns:
                cell.number_format = _PERCENT_FORMAT
    for index, width in enumerate(widths or [], 1):
        sheet.column_dimensions[_column_letter(index)].width = width


def _add_sentiment_chart(sheet: Worksheet, header_row: int, row_count: int) -> None:
    chart = LineChart()
    chart.title = "情绪占比趋势"
    chart.y_axis.title = "占比"
    chart.x_axis.title = "时间窗口（秒）"
    chart.y_axis.scaling.min = 0
    chart.y_axis.scaling.max = 1
    data = Reference(sheet, min_col=8, max_col=10, min_row=header_row, max_row=header_row + row_count)
    categories = Reference(sheet, min_col=1, min_row=header_row + 1, max_row=header_row + row_count)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 22
    sheet.add_chart(chart, "A1")


def _add_topic_chart(sheet: Worksheet, row_count: int) -> None:
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "主题 Top 10"
    chart.x_axis.title = "弹幕数量"
    chart.y_axis.title = "主题"
    data = Reference(sheet, min_col=2, min_row=1, max_row=min(row_count, 10) + 1)
    categories = Reference(sheet, min_col=1, min_row=2, max_row=min(row_count, 10) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.height = 8
    chart.width = 18
    sheet.add_chart(chart, "M1")


def _verify_workbook(path: Path) -> None:
    workbook = load_workbook(path, read_only=True)
    try:
        if workbook.sheetnames != _SHEET_NAMES:
            raise RuntimeError("temporary workbook did not contain the required sheets")
    finally:
        workbook.close()


def _yes_no(value: Any) -> str:
    if isinstance(value, str):
        return "是" if value.strip().lower() == "true" else "否"
    return "是" if value else "否"


def _column_letter(column: int) -> str:
    letters = ""
    while column:
        column, remainder = divmod(column - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters
