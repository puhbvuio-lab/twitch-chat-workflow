from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from twitch_chat_workflow.analysis import AnalysisTables
from twitch_chat_workflow.workbook import write_analysis_workbook


def _tables() -> AnalysisTables:
    return AnalysisTables(
        label_rows=(
            {
                "timestamp_seconds": 1.25,
                "timestamp_ms": 1250,
                "timestamp_iso": "2026-09-21T00:00:01.250Z",
                "author": "alice",
                "text": "清洗后的文字",
                "original_text": "原始弹幕，完整保留。",
                "sentiment": "positive",
                "topic": "直播反馈",
                "interest_signal": True,
                "encoding_warning": False,
                "label_provider": "codex_session",
                "label_model": "test-model",
                "batch_number": 1,
                "message_id": "msg-1",
            },
        ),
        sentiment_rows=(
            {
                "start_seconds": 0,
                "end_seconds": 60,
                "message_count": 2,
                "unique_authors": 2,
                "positive": 1,
                "neutral": 1,
                "negative": 0,
                "positive_share": 0.5,
                "neutral_share": 0.5,
                "negative_share": 0.0,
                "interest_signals": 1,
            },
        ),
        topic_summary_rows=(
            {
                "topic": "直播反馈",
                "message_count": 2,
                "share": 1.0,
                "first_seconds": 1.25,
                "last_seconds": 12.0,
                "positive": 1,
                "neutral": 1,
                "negative": 0,
                "representative_quote_1": "第一条原话",
                "representative_quote_2": "第二条原话",
                "representative_quote_3": "",
            },
        ),
        topic_trend_rows=(
            {
                "start_seconds": 0,
                "end_seconds": 60,
                "topic": "直播反馈",
                "message_count": 2,
                "topic_share": 1.0,
                "positive": 1,
                "neutral": 1,
                "negative": 0,
            },
        ),
        quote_rows=(
            {
                "timestamp_seconds": 1.25,
                "author": "alice",
                "original_text": "原始弹幕，完整保留。",
                "topic": "直播反馈",
                "sentiment": "positive",
                "message_id": "msg-1",
            },
        ),
    )


def test_export_writes_four_readable_sheets_with_charts_and_auditable_original_text(tmp_path: Path) -> None:
    destination = tmp_path / "弹幕分析.xlsx"

    assert write_analysis_workbook(_tables(), destination) == destination

    workbook = openpyxl.load_workbook(destination)
    assert workbook.sheetnames == ["标签明细", "情绪趋势", "主题", "原话"]
    assert workbook["标签明细"].freeze_panes == "A2"
    assert workbook["标签明细"].auto_filter.ref is not None
    assert workbook["标签明细"]["F1"].value == "原始弹幕"
    assert workbook["标签明细"]["F2"].value == "原始弹幕，完整保留。"
    assert workbook["标签明细"]["I2"].value == "是"
    assert workbook["标签明细"]["J2"].value == "否"
    assert workbook["情绪趋势"]["H17"].number_format == "0.0%"
    assert workbook["主题"]["I2"].value == "第一条原话"
    assert workbook["主题"]["C2"].value == 1.0
    assert workbook["主题"]["C2"].number_format == "0.0%"
    assert workbook["主题"]["E6"].value == 1.0
    assert workbook["主题"]["E6"].number_format == "0.0%"
    assert len(workbook["情绪趋势"]._charts) == 1
    assert len(workbook["主题"]._charts) == 1


def test_export_empty_tables_keeps_headers_without_invalid_charts(tmp_path: Path) -> None:
    destination = tmp_path / "弹幕分析.xlsx"
    empty = AnalysisTables(
        label_rows=(), sentiment_rows=(), topic_summary_rows=(), topic_trend_rows=(), quote_rows=()
    )

    write_analysis_workbook(empty, destination)

    workbook = openpyxl.load_workbook(destination)
    assert workbook.sheetnames == ["标签明细", "情绪趋势", "主题", "原话"]
    assert workbook["标签明细"].max_row == 1
    assert workbook["情绪趋势"]["A2"].value == "没有可分析的弹幕"
    assert workbook["主题"]["A2"].value == "没有可分析的弹幕"
    assert not workbook["情绪趋势"]._charts
    assert not workbook["主题"]._charts


def test_export_failure_preserves_existing_workbook_and_removes_only_its_temp_file(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    destination = tmp_path / "弹幕分析.xlsx"
    sentinel = b"previous complete workbook"
    destination.write_bytes(sentinel)

    def fail_save(self: openpyxl.Workbook, _path: str) -> None:
        raise OSError("disk full")

    monkeypatch.setattr(openpyxl.Workbook, "save", fail_save)

    with pytest.raises(OSError, match="disk full"):
        write_analysis_workbook(_tables(), destination)

    assert destination.read_bytes() == sentinel
    assert not list(tmp_path.glob(".弹幕分析.*.tmp.xlsx"))
